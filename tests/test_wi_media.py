import csv
from pathlib import Path
import subprocess
import tempfile
import unittest
from unittest.mock import patch
from PIL import Image
from openpyxl import load_workbook
from lynx_bulk import wi_rows, build_table, default_fields, write_excel
from lynx_wi_media import acquire_wi, acquired_rows, sources_for
from lynx_tasks import TaskContext, TaskCancelled
from wi_factory import make_wi_zip


class WIMediaTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.task = TaskContext()
        self.zip = make_wi_zip(self.root)
        self.rows, _ = wi_rows(self.task, self.zip)

    def test_zip_local_photos_to_excel_preserves_metadata_and_actual_names(self):
        Image.new('RGB', (4, 4)).save(self.root / 'LYNX.JPG')
        result = acquire_wi(self.task, sources_for(self.rows), self.root)
        self.assertEqual(result.failed, ['gs://bucket/fox.png'])
        rows, missing = acquired_rows(self.task, self.rows, result.available, True)
        self.assertEqual(missing, result.failed)
        self.assertEqual(rows[0]['metadata']['projects.project_name'], ['Synthetic project'])
        table = build_table(rows, default_fields())
        self.assertEqual(table[0]['Encounter.mediaAsset0'], 'LYNX.JPG')
        output = self.root / 'out.xlsx'
        write_excel(self.task, table, output)
        book = load_workbook(output)
        self.assertEqual(book.active.max_row, 2)
        book.close()
        (self.root / 'LYNX.JPG').unlink()
        self.assertEqual(acquired_rows(self.task, self.rows, result.available)[0], [])
        with self.assertRaises(ValueError):
            build_table(rows, default_fields())

    def test_local_duplicates_and_invalid_images_are_not_exported(self):
        (self.root / 'nested').mkdir()
        for path in (self.root / 'lynx.jpg', self.root / 'nested' / 'LYNX.JPG'):
            Image.new('RGB', (4, 4)).save(path)
        (self.root / 'fox.png').write_text('invalid image')
        result = acquire_wi(self.task, sources_for(self.rows), self.root)
        self.assertEqual(len(result.failed), 2)
        self.assertFalse(result.available)

    def test_source_name_collision_is_rejected(self):
        with self.assertRaises(ValueError):
            sources_for([{'media': ['gs://one/a.jpg', 'gs://two/A.JPG']}])

    def test_gsutil_arguments_failure_retry_and_manifest(self):
        sources = sources_for(self.rows)
        def download(args, **kwargs):
            self.assertEqual(args[:2], ['/fake/gsutil', 'cp'])
            self.assertTrue(kwargs['check'])
            if args[2].endswith('lynx.jpg'):
                raise subprocess.CalledProcessError(1, args, stderr='access denied')
            Image.new('RGB', (4, 4)).save(args[3])
        with patch('lynx_wi_media.subprocess.run', side_effect=download) as run:
            first = acquire_wi(self.task, sources, self.root, '/fake/gsutil')
        self.assertEqual(run.call_count, 2)
        self.assertEqual(first.failed, ['gs://bucket/lynx.jpg'])
        with open(first.manifest) as stream:
            manifest = list(csv.DictReader(stream))
        self.assertEqual({row['status'] for row in manifest}, {'completed', 'failed'})
        with patch('lynx_wi_media.subprocess.run', side_effect=lambda args, **kw: Image.new('RGB', (4, 4)).save(args[3])):
            retry = acquire_wi(self.task, first.failed, self.root, '/fake/gsutil')
        first.available.update(retry.available)
        rows, missing = acquired_rows(self.task, self.rows, first.available)
        self.assertFalse(missing)
        self.assertEqual(len(build_table(rows, default_fields())), 2)
        self.assertTrue(Path(first.available['gs://bucket/fox.png']).is_file())

    def test_no_wildcards_or_invalid_downloads_committed(self):
        with self.assertRaises(ValueError):
            acquire_wi(self.task, ['gs://bucket/manifest.csv'], self.root, '/fake/gsutil')
        with patch('lynx_wi_media.subprocess.run') as run:
            result = acquire_wi(self.task, ['gs://bucket/*.jpg'], self.root, '/fake/gsutil')
        run.assert_not_called()
        self.assertEqual(len(result.failed), 1)
        with patch('lynx_wi_media.subprocess.run', side_effect=lambda args, **kw: Path(args[3]).write_text('not a photo')):
            result = acquire_wi(self.task, ['gs://bucket/a.jpg'], self.root, '/fake/gsutil')
        self.assertFalse(result.available)
        self.assertFalse((Path(result.manifest).parent / 'a.jpg').exists())

    def test_cancellation_preserves_completed_download_and_manifest(self):
        def download(args, **kwargs):
            Image.new('RGB', (4, 4)).save(args[3])
            self.task.stop.set()
        with patch('lynx_wi_media.subprocess.run', side_effect=download):
            with self.assertRaises(TaskCancelled):
                acquire_wi(self.task, sources_for(self.rows), self.root, '/fake/gsutil')
        manifest = next(self.root.rglob('manifest.csv'))
        with manifest.open() as stream:
            rows = list(csv.DictReader(stream))
        self.assertEqual(rows[0]['status'], 'completed')
        self.assertTrue((manifest.parent / rows[0]['file']).is_file())
