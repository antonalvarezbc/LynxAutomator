import tempfile
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from camtrap_factory import make_package
from lynx_bulk import (default_fields, save_profile, load_profile, normalize, build_table,
                       write_excel, wi_rows, camtrap_rows)
from lynx_camtrap import read_package, select_media, acquire_media
from lynx_tasks import TaskContext


class BulkTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.task = TaskContext()

    def row(self):
        photo = self.root / 'actual.png'
        photo.write_bytes(b'test')
        return normalize('Lynx pardinus', '2024-06-01T23:30:00-03:00', photo, 'd', 'e', 'p', locality='test')

    def test_profile_and_excel_literal(self):
        fields = default_fields()
        fields.append(dict(name='Encounter.remarks', source='fixed', value='=1+2', type='text', enabled=True))
        profile = self.root / 'settings' / 'profile.json'
        save_profile(fields, profile)
        self.assertEqual(load_profile(profile), fields)
        table = build_table([self.row()], fields)
        self.assertEqual(table[0]['Encounter.hour'], 23)
        self.assertEqual(table[0]['Encounter.mediaAsset0'], 'actual.png')
        path = self.root / 'out.xlsx'
        write_excel(self.task, table, path)
        book = load_workbook(path)
        sheet = book.active
        column = [c.value for c in sheet[1]].index('Encounter.remarks') + 1
        self.assertEqual(sheet.cell(2, column).data_type, 's')
        self.assertEqual(sheet.cell(2, column).value, '=1+2')
        book.close()

    def test_invalid_required_fields_and_duplicates(self):
        fields = default_fields()
        fields[0]['enabled'] = False
        with self.assertRaises(ValueError):
            build_table([self.row()], fields)
        with self.assertRaises(ValueError):
            build_table([self.row()], default_fields() * 2)
        row = self.row()
        row['locality'] = ''
        with self.assertRaises(ValueError):
            build_table([row], default_fields())

    def test_custom_types_and_deleted_media(self):
        fields = default_fields() + [dict(name='Encounter.custom', source='fixed', value='3.5', type='integer')]
        with self.assertRaises(ValueError):
            build_table([self.row()], fields)
        fields[-1]['type'] = 'decimal'
        row = self.row()
        self.assertEqual(build_table([row], fields)[0]['Encounter.custom'], 3.5)
        Path(row['media'][0]).unlink()
        with self.assertRaises(ValueError):
            build_table([row], fields)

    def test_camtrap_downloads_and_retries(self):
        package = read_package(self.task, make_package(self.root / 'package') / 'datapackage.json')
        package.deployments[0]['locationName'] = 'test'
        package.deployments[0]['cameraModel'] = 'Model X'
        records, _ = select_media(self.task, package, {'Lynx pardinus'})
        result = acquire_media(self.task, package, records, self.root, local_only=True)
        rows, missing = camtrap_rows(self.task, package, records, [result.output_directory])
        self.assertEqual(len(rows), 2)
        self.assertEqual(missing, ['m2'])
        self.assertEqual({r['species'] for r in rows}, {'Lynx pardinus'})
        result2 = acquire_media(self.task, package, records, self.root, local_only=True)
        retried, _ = camtrap_rows(self.task, package, records, [result.output_directory, result2.output_directory])
        self.assertEqual(len(retried), 2)
        table = build_table(retried, default_fields() + [dict(name='Sighting.comments', source='template', value='Model: {deployment.cameraModel}', type='text')])
        self.assertTrue(all(r['Sighting.comments'] == 'Model: Model X' for r in table))
        self.assertTrue(all(r['Encounter.mediaAsset0'].endswith('.jpg') for r in table))

    def test_wi_species_and_events(self):
        images = []
        for i, (species, second, count) in enumerate([('Lynx pardinus', 0, 1), ('Lynx pardinus', 2, 1),
                                                    ('Vulpes vulpes', 3, 1), ('Lynx pardinus', 4, 2)]):
            images.append(dict(project_id='p', deployment_id='d', location=f'gs://bucket/{i}.png',
                               timestamp=f'2024-01-01T00:00:0{second}Z', scientific_name=species,
                               number_of_objects=count))
        pd.DataFrame(images).to_csv(self.root / 'images.csv', index=False)
        pd.DataFrame([dict(project_id='p', deployment_id='d', placename='test', latitude=0, longitude=0)]).to_csv(self.root / 'deployments.csv', index=False)
        rows, missing = wi_rows(self.task, self.root / 'images.csv', self.root / 'deployments.csv', group=True, threshold=3)
        self.assertFalse(missing)
        self.assertEqual(len(rows), 3)
        self.assertEqual(len(rows[0]['media']), 2)
        self.assertEqual(len({r['eventID'] for r in rows}), 3)
        self.assertEqual(len(build_table(rows, default_fields())), 3)

    def test_camtrap_event_overlap_deduplicates_and_orders_dates(self):
        package = read_package(self.task, make_package(self.root / 'package') / 'datapackage.json')
        package.deployments[0]['locationName'] = 'test'
        package.deployments[0]['cameraModel'] = 'Model X'
        package.observations.append(dict(observationID='extra', deploymentID='d1', observationType='animal',
                                         observationLevel='event', scientificName='Lynx pardinus',
                                         eventID='all', eventStart='2024-01-01T00:00:00Z', eventEnd='2024-01-01T00:01:00Z'))
        package.observations[2]['eventID'] = 'all'
        records, _ = select_media(self.task, package, {'Lynx pardinus'})
        result = acquire_media(self.task, package, records, self.root, local_only=True)
        rows, _ = camtrap_rows(self.task, package, list(reversed(records)), [result.output_directory])
        self.assertEqual(len(rows), 1)
        self.assertEqual(len(rows[0]['media']), 2)
        self.assertEqual(rows[0]['minutes'], 0)

    def test_wi_zip_matches_csv_without_photos(self):
        import zipfile
        self.test_wi_species_and_events()
        archive = self.root / 'export.zip'
        with zipfile.ZipFile(archive, 'w') as z:
            for name in ['images.csv', 'deployments.csv']:
                z.write(self.root / name, 'export/' + name)
        rows, missing = wi_rows(self.task, archive)
        self.assertFalse(missing)
        self.assertEqual(len(rows), 4)
        self.assertEqual(build_table(rows, default_fields())[0]['Encounter.mediaAsset0'], '0.png')
        with zipfile.ZipFile(archive, 'a') as z:
            z.write(self.root / 'images.csv', 'other/images.csv')
        with self.assertRaisesRegex(ValueError, 'único conjunto'):
            wi_rows(self.task, archive)

    def test_named_profiles_migrate_and_remain_independent(self):
        import json
        from lynx_bulk import read_profiles
        path = self.root / 'profiles.json'
        fields = default_fields()
        path.write_text(json.dumps({'version': 1, 'fields': fields}))
        save_profile(fields, path, name='Doñana')
        fields[-1]['value'] = 'another user'
        save_profile(fields, path, name='Sierra Morena')
        self.assertEqual(set(read_profiles(path)), {'Default', 'Doñana', 'Sierra Morena'})
        self.assertEqual(load_profile(path, name='Doñana')[-1]['value'], '')
        self.assertEqual(load_profile(path, name='Sierra Morena')[-1]['value'], 'another user')

    def test_wi_project_suffix_and_multiple_parts(self):
        import zipfile
        self.test_wi_species_and_events()
        archive = self.root / 'project.zip'
        with zipfile.ZipFile(archive, 'w') as z:
            z.write(self.root / 'images.csv', 'export/images_2001260.csv')
            z.write(self.root / 'deployments.csv', 'export/deployments.csv')
        rows, _ = wi_rows(self.task, archive)
        self.assertEqual(len(build_table(rows, default_fields())), 4)

    def test_comments_preserve_grouped_metadata_and_profile(self):
        from lynx_bulk import attach_metadata, merge_metadata
        first, second = self.row(), self.row()
        attach_metadata(first, deployment=[{'setupBy': 'Alice', 'cameraID': 'C1', 'cameraModel': 'Model X'}],
                        observation=[{'comments': 'first'}])
        attach_metadata(second, deployment=[{'setupBy': 'Alice', 'cameraID': 'C1', 'cameraModel': 'Model X'}],
                        observation=[{'comments': 'second {literal}'}])
        merge_metadata(first, second)
        field = dict(name='Sighting.comments', source='template', type='text', enabled=True,
                     value='Camera: {deployment.cameraID}; Model: {deployment.cameraModel}; Setup: {deployment.setupBy}; Notes: {observation.comments}')
        fields = default_fields() + [field]
        profile = self.root / 'profile.json'
        save_profile(fields, profile, 'Camera notes')
        table = build_table([first], load_profile(profile, 'Camera notes'))
        self.assertEqual(table[0]['Sighting.comments'], 'Camera: C1; Model: Model X; Setup: Alice; Notes: first | second {literal}')
        self.assertIn('Encounter.sightingID', table[0])
        field['value'] = '{deployment.typo}'
        with self.assertRaisesRegex(ValueError, 'deployment.typo'):
            build_table([first], default_fields() + [field])

    def test_required_warning_names_and_media_subfields(self):
        from lynx_bulk import required_warnings
        warnings = required_warnings({'Encounter.decimalLatitude': 0, 'Encounter.decimalLongitude': 0})
        self.assertEqual(len(warnings), 4)
        self.assertIn('Falta Encounter.year', warnings)
        fields = default_fields() + [dict(name='Encounter.mediaAsset0.keywords', source='fixed', value='camera', type='text')]
        self.assertEqual(build_table([self.row()], fields)[0]['Encounter.mediaAsset0.keywords'], 'camera')
        fields[0]['enabled'] = False
        with self.assertRaisesRegex(ValueError, 'Encounter.genus'):
            build_table([self.row()], fields)

    def test_local_wildlife_example_if_available(self):
        fixture = Path(__file__).parent / 'fixtures' / 'Wildlife Insights'
        archives = list(fixture.glob('*.zip'))
        if not archives:
            self.skipTest('Optional private local fixture')
        rows, missing = wi_rows(self.task, archives[0])
        table = build_table(rows, default_fields())
        self.assertGreater(len(table), 0)
        self.assertFalse(missing)
        output = self.root / 'real-example.xlsx'
        write_excel(self.task, table, output)
        book = load_workbook(output)
        self.assertEqual(book.active.max_row, len(table) + 1)
        book.close()

    def test_additional_projects_and_cameras_are_joined_without_cross_project_leaks(self):
        import zipfile
        self.test_wi_species_and_events()
        pd.DataFrame([dict(project_id='p', project_name='Correct'), dict(project_id='other', project_name='Wrong')]).to_csv(self.root / 'projects.csv', index=False)
        pd.DataFrame([dict(deployment_id='d', camera_model='Model A')]).to_csv(self.root / 'cameras.csv', index=False)
        archive = self.root / 'metadata.zip'
        with zipfile.ZipFile(archive, 'w') as z:
            for name in ['images.csv', 'deployments.csv', 'projects.csv', 'cameras.csv']:
                z.write(self.root / name, 'export/' + name)
        rows, _ = wi_rows(self.task, archive)
        fields = default_fields() + [dict(name='Sighting.comments', source='template', type='text', value='{projects.project_name}; {cameras.camera_model}')]
        table = build_table(rows, fields)
        self.assertTrue(all(r['Sighting.comments'] == 'Correct; Model A' for r in table))
        rows, _ = wi_rows(self.task, self.root / 'images.csv', self.root / 'deployments.csv', extra_paths=[self.root / 'projects.csv', self.root / 'cameras.csv'])
        self.assertEqual(build_table(rows, fields), table)

    def test_shared_grouping_interval_and_explicit_events(self):
        from lynx_bulk import group_rows, attach_metadata
        rows = []
        for i in range(3):
            row = self.row()
            row['timestamp'] = f'2024-01-01T00:00:0{i * 2}Z'
            row['media'] = [str(i)]
            attach_metadata(row, deployment=[{'deploymentID': 'd'}])
            rows.append(row)
        self.assertEqual(len(group_rows(self.task, rows, 1)), 3)
        self.assertEqual(len(group_rows(self.task, rows, 3)), 1)
        rows[0]['_event'] = 'event1'
        rows[2]['_event'] = 'event1'
        result = group_rows(self.task, rows, 0)
        self.assertEqual(len(result), 2)
        self.assertEqual(len(result[0]['media']), 2)

    def test_wi_filter_excludes_unselected_species(self):
        from lynx_bulk import wi_species
        self.test_wi_species_and_events()
        names = wi_species(self.task, self.root / 'images.csv', self.root / 'deployments.csv')
        self.assertEqual(set(names), {'Lynx pardinus', 'Vulpes vulpes'})
        rows, _ = wi_rows(self.task, self.root / 'images.csv', self.root / 'deployments.csv', species={'Vulpes vulpes'})
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]['species'], 'Vulpes vulpes')

    def test_folder_adapter_missing_dates_and_explicit_identity(self):
        from lynx_folder_bulk import folder_rows
        from PIL import Image
        Image.new('RGB', (8, 8)).save(self.root / 'Animal1 left.jpg')
        rows, missing = folder_rows(self.task, self.root, 'Lynx pardinus')
        self.assertFalse(rows)
        self.assertEqual(len(missing), 1)
        rows, missing = folder_rows(self.task, self.root, 'Lynx pardinus', fallback_year=2020)
        self.assertFalse(missing)
        self.assertEqual(rows[0]['year'], 2020)
        self.assertEqual(rows[0]['month'], '')
        self.assertEqual(rows[0]['individualID'], '')
        rows, _ = folder_rows(self.task, self.root, 'Lynx pardinus', fallback_year=2020, identity='filename')
        self.assertEqual(rows[0]['individualID'], 'Animal1')

    def test_catalog_without_dates_exports_blank_time(self):
        from lynx_folder_bulk import folder_rows
        from PIL import Image
        for name in ('one.jpg', 'two.jpg'):
            Image.new('RGB', (2, 2)).save(self.root / name)
        rows, missing = folder_rows(self.task, self.root, 'Lynx pardinus',
                                    allow_undated=True, group=True)
        self.assertEqual(len(rows), 2)
        self.assertEqual(missing, [])
        for row in rows:
            row['locality'] = 'Catalog'
            self.assertEqual(row['timestamp'], '')
        table = build_table(rows, default_fields())
        path = self.root / 'catalog.xlsx'
        write_excel(self.task, table, path)
        book = load_workbook(path)
        headers = [cell.value for cell in book.active[1]]
        for name in ('year', 'month', 'day', 'hour', 'minutes'):
            column = headers.index('Encounter.' + name) + 1
            self.assertIsNone(book.active.cell(2, column).value)
        book.close()
